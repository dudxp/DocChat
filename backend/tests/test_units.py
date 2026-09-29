"""Testes que não precisam de banco: chunking, citações, métricas e o juiz."""

import io
from types import SimpleNamespace

import pytest
from conftest import MANUAL

from app.rag.evaluation import citation_hit, parse_judge, retrieval_metrics, summarize, token_f1
from app.rag.extraction import IngestionError, Unit, clean_text, extract
from app.rag.ingestion import split_units
from app.rag.pipeline import Answer, extract_citations, format_context
from app.rag.providers import HashingEmbeddings
from app.rag.retrieval import RetrievedChunk, build_or_tsquery, reciprocal_rank_fusion


def chunk(unit: int, document_id: int = 1) -> RetrievedChunk:
    return RetrievedChunk(
        chunk_id=unit,
        document_id=document_id,
        filename="a.pdf",
        kind="pdf",
        unit=unit,
        location=f"página {unit}",
        content="x",
        similarity=0.5,
        score=0.1,
    )


def units(*texts: str) -> list[Unit]:
    return [Unit(ordinal=i, location=f"página {i}", text=t) for i, t in enumerate(texts, start=1)]


class TestIngestion:
    def test_pdf_has_one_unit_per_page(self):
        kind, extracted = extract("manual.pdf", MANUAL.read_bytes())
        assert kind == "pdf"
        assert len(extracted) == 6
        assert "1,5 kW" in extracted[1].text
        assert extracted[1].location == "página 2"

    def test_chunks_never_cross_units(self):
        chunks = split_units(units("a " * 800, "", "b " * 300), chunk_size=500, chunk_overlap=50)
        assert {c.unit for c in chunks} == {1, 3}
        assert all(set(c.content.split()) == {"a"} for c in chunks if c.unit == 1)
        assert [c.index for c in chunks] == list(range(len(chunks)))

    def test_chunk_inherits_the_location_of_its_unit(self):
        chunks = split_units([Unit(ordinal=2, location="planilha Custos", text="a " * 100)], 500, 50)
        assert {c.location for c in chunks} == {"planilha Custos"}
        assert {c.unit for c in chunks} == {2}

    def test_clean_text_rejoins_hyphenated_words_and_drops_nul(self):
        assert clean_text("manu-\ntenção\x00  preventiva") == "manutenção preventiva"


class TestExtraction:
    """Cada formato se divide de um jeito, e é essa divisão que a citação vai endereçar."""

    def test_word_splits_on_headings_and_keeps_tables_in_place(self):
        from docx import Document as Docx

        document = Docx()
        document.add_heading("Garantia", level=1)
        document.add_paragraph("O prazo é de 12 meses.")
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Peça"
        table.rows[0].cells[1].text = "Parafuso"
        document.add_heading("Manutenção", level=2)
        document.add_paragraph("Lubrificar a cada 500 horas.")
        buffer = io.BytesIO()
        document.save(buffer)

        kind, extracted = extract("manual.docx", buffer.getvalue())
        assert kind == "docx"
        assert [u.location for u in extracted] == ["Garantia", "Manutenção"]
        # A tabela está na seção em que foi escrita, não empurrada para o fim do documento.
        assert "Parafuso" in extracted[0].text
        assert "500 horas" in extracted[1].text

    def test_spreadsheet_has_one_unit_per_sheet_named_after_it(self):
        from openpyxl import Workbook

        workbook = Workbook()
        workbook.active.title = "Custos"
        workbook.active.append(["Item", "Valor"])
        workbook.active.append(["Motor", 1500])
        workbook.create_sheet("Prazos").append(["Entrega", "30 dias"])
        buffer = io.BytesIO()
        workbook.save(buffer)

        kind, extracted = extract("planilha.xlsx", buffer.getvalue())
        assert kind == "xlsx"
        assert [u.location for u in extracted] == ["planilha Custos", "planilha Prazos"]
        assert "Motor | 1500" in extracted[0].text

    def test_presentation_has_one_unit_per_slide(self):
        from pptx import Presentation

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = "Resultados do trimestre"
        buffer = io.BytesIO()
        presentation.save(buffer)

        kind, extracted = extract("apresentacao.pptx", buffer.getvalue())
        assert kind == "pptx"
        assert [u.location for u in extracted] == ["slide 1"]
        assert "Resultados do trimestre" in extracted[0].text

    def test_markdown_splits_on_headings(self):
        source = b"# Introducao\ntexto um\n\n## Detalhes\ntexto dois\n"
        kind, extracted = extract("notas.md", source)
        assert kind == "text"
        assert [u.location for u in extracted] == ["Introducao", "Detalhes"]

    def test_csv_row_carries_its_column_names(self):
        """Sem o cabeçalho junto, o trecho recuperado é uma linha de números sem significado."""
        _, extracted = extract("dados.csv", b"item,valor\nmotor,1500\n")
        assert extracted[0].text == "item: motor | valor: 1500"
        assert extracted[0].location == "planilha"

    def test_plain_text_saved_on_windows_is_not_mangled(self):
        _, extracted = extract("nota.txt", "manutenção preventiva".encode("cp1252"))
        assert extracted[0].text == "manutenção preventiva"

    def test_unsupported_extension_is_rejected_with_the_accepted_list(self):
        with pytest.raises(IngestionError, match="Formato não suportado"):
            extract("foto.jpg", b"\xff\xd8\xff")

    def test_empty_document_yields_no_units(self):
        assert extract("vazio.txt", b"   \n  ")[1] == []


class TestCitations:
    def test_extracts_unique_citations_in_order(self):
        assert extract_citations("A [2]. B [1][2]. C [2]", 3) == [2, 1]

    def test_ignores_numbers_outside_the_sources(self):
        assert extract_citations("Veja [0] e [7] e [3]", 3) == [3]

    def test_context_numbers_sources_from_one(self):
        context = format_context([chunk(4), chunk(9)])
        assert context.startswith("[1] (documento: a.pdf, página 4)")
        assert "[2] (documento: a.pdf, página 9)" in context


class TestRetrievalHelpers:
    def test_tsquery_uses_or_between_terms(self):
        assert build_or_tsquery("Qual a pressão máxima?") == "máxima | pressão | qual"

    def test_tsquery_of_punctuation_only_is_none(self):
        assert build_or_tsquery("?? !") is None

    def test_rrf_rewards_items_present_in_both_rankings(self):
        scores = reciprocal_rank_fusion([[1, 2, 3], [3, 4]])
        assert max(scores, key=scores.get) == 3

    def test_hashing_embeddings_are_normalized_and_deterministic(self):
        emb = HashingEmbeddings(64)
        a, b = emb.embed_query("motor trifásico"), emb.embed_query("motor trifásico")
        assert a == b
        assert abs(sum(v * v for v in a) - 1.0) < 1e-9


class TestMetrics:
    case = SimpleNamespace(document_id=1, expected_unit=5)

    def answer(self, units: list[int], cited: list[int]) -> Answer:
        return Answer(question="q", search_query="q", answer="", sources=[chunk(u) for u in units], cited=cited)

    def test_reciprocal_rank_uses_position_of_first_correct_chunk(self):
        assert retrieval_metrics(self.answer([2, 5, 5], []), self.case) == (True, 0.5)

    def test_miss_scores_zero(self):
        assert retrieval_metrics(self.answer([1, 2], []), self.case) == (False, 0.0)

    def test_same_page_in_another_document_is_not_a_hit(self):
        answer = Answer(question="q", search_query="q", answer="", sources=[chunk(5, document_id=2)])
        assert retrieval_metrics(answer, self.case) == (False, 0.0)

    def test_case_without_unit_is_not_measured(self):
        case = SimpleNamespace(document_id=None, expected_unit=None)
        assert retrieval_metrics(self.answer([1], []), case) == (None, None)
        assert citation_hit(self.answer([1], [1]), case) is None

    def test_citation_hit_only_when_the_cited_source_is_the_right_page(self):
        assert citation_hit(self.answer([5, 2], [1]), self.case) is True
        assert citation_hit(self.answer([5, 2], [2]), self.case) is False

    def test_token_f1_ignores_accents_and_citation_marks(self):
        assert token_f1("Óleo sintético ISO VG 220 [1]", "oleo sintetico ISO VG 220") == 1.0
        assert token_f1("nada a ver", "óleo sintético") == 0.0

    def test_summary_skips_unmeasured_values(self):
        results = [
            SimpleNamespace(
                retrieval_hit=True,
                reciprocal_rank=1.0,
                citation_hit=True,
                answer_f1=0.5,
                judge_score=None,
                latency_ms=100,
            ),
            SimpleNamespace(
                retrieval_hit=None,
                reciprocal_rank=None,
                citation_hit=None,
                answer_f1=0.3,
                judge_score=None,
                latency_ms=300,
            ),
        ]
        summary = summarize(results)
        assert summary["hit_rate"] == 1.0
        assert summary["answer_f1"] == 0.4
        assert summary["judge_score"] is None
        assert summary["avg_latency_ms"] == 200


class TestJudgeParsing:
    def test_reads_json_even_with_text_around_it(self):
        assert parse_judge('Claro! {"score": 0.5, "reason": "incompleta"}') == (0.5, "incompleta")

    def test_clamps_score_to_zero_one(self):
        assert parse_judge('{"score": 3}')[0] == 1.0

    def test_unparseable_output_has_no_score(self):
        assert parse_judge("nota oito")[0] is None
