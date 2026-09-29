"""Extração do texto por formato.

Cada formato responde a mesma pergunta de um jeito diferente: **em que o documento se divide?**
No PDF é a página, na planilha é a aba, na apresentação é o slide, no texto é a seção. Essa divisão
é o que a citação endereça, então ela nasce aqui e atravessa o resto do sistema como `Unit`.
"""

import csv
import io
import re
from collections.abc import Callable, Iterator
from dataclasses import dataclass
from pathlib import Path

from app.domain.models import DocumentKind

MAX_LOCATION = 80


class IngestionError(ValueError):
    pass


@dataclass(frozen=True)
class Unit:
    """Uma divisão do documento: o ordinal que endereça e o rótulo que a pessoa lê."""

    ordinal: int
    location: str
    text: str


def clean_text(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"-\n(\w)", r"\1", text)  # palavra hifenizada na quebra de linha
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _label(text: str, fallback: str) -> str:
    """O rótulo vem do próprio documento quando ele tem um título; senão, do ordinal."""
    first = " ".join(text.split())[:MAX_LOCATION].strip()
    return first or fallback


def _pdf(data: bytes) -> list[Unit]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [clean_text(page.extract_text() or "") for page in reader.pages]
    except (PdfReadError, ValueError, KeyError) as exc:
        raise IngestionError("Não foi possível ler o PDF. O arquivo pode estar corrompido.") from exc
    return [Unit(ordinal=i, location=f"página {i}", text=t) for i, t in enumerate(pages, start=1)]


def _docx_blocks(document) -> Iterator[tuple[str, str]]:
    """Percorre parágrafos e tabelas na ordem em que aparecem no corpo.

    O python-docx expõe as duas coleções separadas, e usá-las assim desmonta a ordem original:
    uma tabela no meio do texto iria parar no fim e mudaria de seção.
    """
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            paragraph = Paragraph(child, document)
            yield paragraph.style.name or "", paragraph.text
        elif child.tag.endswith("}tbl"):
            rows = [" | ".join(c.text.strip() for c in row.cells) for row in Table(child, document).rows]
            yield "", "\n".join(r for r in rows if r.strip(" |"))


def _docx(data: bytes) -> list[Unit]:
    from docx import Document as DocxDocument
    from docx.opc.exceptions import PackageNotFoundError

    try:
        document = DocxDocument(io.BytesIO(data))
    except (PackageNotFoundError, KeyError, ValueError) as exc:
        raise IngestionError("Não foi possível ler o documento do Word. O arquivo pode estar corrompido.") from exc

    units: list[Unit] = []
    location, buffer = "início do documento", []

    def flush() -> None:
        text = clean_text("\n".join(buffer))
        if text:
            units.append(Unit(ordinal=len(units) + 1, location=location, text=text))

    for style, text in _docx_blocks(document):
        if style.startswith("Heading") and text.strip():
            flush()
            buffer = [text]
            location = _label(text, f"seção {len(units) + 1}")
        elif text.strip():
            buffer.append(text)
    flush()
    return units


def _xlsx(data: bytes) -> list[Unit]:
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        # data_only: interessa o valor calculado da fórmula, não a fórmula em si.
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (InvalidFileException, KeyError, ValueError) as exc:
        raise IngestionError("Não foi possível ler a planilha. O arquivo pode estar corrompido.") from exc

    units: list[Unit] = []
    for sheet in workbook.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                rows.append(" | ".join(cells))
        text = clean_text("\n".join(rows))
        if text:
            units.append(Unit(ordinal=len(units) + 1, location=f"planilha {sheet.title}", text=text))
    workbook.close()
    return units


def _pptx(data: bytes) -> list[Unit]:
    from pptx import Presentation
    from pptx.exc import PackageNotFoundError

    try:
        presentation = Presentation(io.BytesIO(data))
    except (PackageNotFoundError, KeyError, ValueError) as exc:
        raise IngestionError("Não foi possível ler a apresentação. O arquivo pode estar corrompido.") from exc

    units: list[Unit] = []
    for number, slide in enumerate(presentation.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            elif shape.has_table:
                parts += [" | ".join(c.text.strip() for c in row.cells) for row in shape.table.rows]
        text = clean_text("\n".join(parts))
        if text:
            units.append(Unit(ordinal=number, location=f"slide {number}", text=text))
    return units


_HEADING = re.compile(r"^#{1,6}\s+(.+)$")


def _markdown(data: bytes) -> list[Unit]:
    units: list[Unit] = []
    location, buffer = "início do documento", []

    def flush() -> None:
        text = clean_text("\n".join(buffer))
        if text:
            units.append(Unit(ordinal=len(units) + 1, location=location, text=text))

    for line in _decode(data).splitlines():
        heading = _HEADING.match(line)
        if heading:
            flush()
            buffer = [line]
            location = _label(heading.group(1), f"seção {len(units) + 1}")
        else:
            buffer.append(line)
    flush()
    return units


def _csv(data: bytes) -> list[Unit]:
    """Cada linha vira 'coluna: valor', senão o trecho perde o cabeçalho e não se explica sozinho."""
    reader = csv.reader(io.StringIO(_decode(data)))
    rows = [r for r in reader if any(c.strip() for c in r)]
    if not rows:
        return []
    header, *body = rows

    def described(row: list[str]) -> str:
        pairs = zip(header, row, strict=False)
        return " | ".join(f"{name.strip()}: {value.strip()}" for name, value in pairs if value.strip())

    lines = [described(row) for row in body]
    text = clean_text("\n".join(lines) if lines else " | ".join(header))
    return [Unit(ordinal=1, location="planilha", text=text)] if text else []


def _plain(data: bytes) -> list[Unit]:
    text = clean_text(_decode(data))
    return [Unit(ordinal=1, location="documento", text=text)] if text else []


def _decode(data: bytes) -> str:
    """Arquivo de texto salvo no Windows costuma vir em cp1252, e utf-8 estrito quebraria nele."""
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


EXTRACTORS: dict[str, tuple[DocumentKind, Callable[[bytes], list[Unit]]]] = {
    ".pdf": (DocumentKind.PDF, _pdf),
    ".docx": (DocumentKind.DOCX, _docx),
    ".xlsx": (DocumentKind.XLSX, _xlsx),
    ".pptx": (DocumentKind.PPTX, _pptx),
    ".md": (DocumentKind.TEXT, _markdown),
    ".markdown": (DocumentKind.TEXT, _markdown),
    ".csv": (DocumentKind.TEXT, _csv),
    ".txt": (DocumentKind.TEXT, _plain),
}

SUPPORTED_EXTENSIONS = tuple(sorted(EXTRACTORS))


def extract(filename: str, data: bytes) -> tuple[DocumentKind, list[Unit]]:
    extension = Path(filename).suffix.lower()
    if extension not in EXTRACTORS:
        accepted = ", ".join(SUPPORTED_EXTENSIONS)
        raise IngestionError(f"Formato não suportado. Aceitos: {accepted}.")
    kind, extractor = EXTRACTORS[extension]
    return kind, extractor(data)
